#!/usr/bin/env python3
"""
DocuSign monthly envelope usage extractor
=========================================

Produces ONE Excel file with two sheets:

  Sheet "MonthlyUsage"  - one row per calendar month, envelope counts split by
                          status, plus date columns shaped for Power BI filters.
  Sheet "TotalsToDate"  - a single-row rollup of everything in MonthlyUsage,
                          rewritten on every run.

Incremental by design
---------------------
On each run the script reads the existing workbook and only calls DocuSign for
months it does not already have. Two exceptions, both deliberate:

  * The current (in-flight) month is ALWAYS re-fetched, because last time it ran
    it only covered part of the month.
  * Any stored row flagged IsComplete = FALSE is re-fetched for the same reason.

So a completed past month is fetched exactly once, ever.

Partial months
--------------
For the current month the window runs from the 1st to TODAY. DaysCounted tells
you how many days that row actually covers and DaysInMonth tells you how many
the month has, so you can normalise in Power BI (e.g. a run-rate measure of
Total / DaysCounted * DaysInMonth) instead of drawing a misleading short bar.

Auth (unchanged from the previous script)
-----------------------------------------
  A) set DS_ACCESS_TOKEN=<token>                 -- used as-is, nothing rotated
  B) set DS_INTEGRATION_KEY / DS_CLIENT_SECRET / DS_REFRESH_TOKEN
                                                 -- mints a token, ROTATES the
                                                    refresh token

Also set
--------
  set DS_ACCOUNT_ID=<API account ID>
  set DS_START_DATE=2025-05-01        (contract / context start; day is ignored,
                                       the month is always taken whole)

Optional
--------
  set DS_OUTPUT_FILE=docusign_envelope_usage.xlsx
  set DS_COUNT_MODE=fast              fast | full   (default fast)
  set DS_TZ_OFFSET=+00:00             account timezone, e.g. +05:30 or -05:00
  set DS_FROM_TO_STATUS=created       created | changed
  set DS_REBUILD=1                    ignore existing rows, refetch everything

Requirements
------------
    pip install requests openpyxl
"""

import base64
import calendar
import os
import sys
import time
from datetime import datetime, timedelta, timezone
from pathlib import Path

import requests

try:
    from openpyxl import Workbook, load_workbook
except ImportError:
    sys.exit("openpyxl is missing.  Run:  pip install openpyxl")

# ---------------------------------------------------------------------------
# CONFIG
# ---------------------------------------------------------------------------

ENVIRONMENT = os.getenv("DS_ENV", "production")

ACCESS_TOKEN = os.getenv("DS_ACCESS_TOKEN", "")
INTEGRATION_KEY = os.getenv("DS_INTEGRATION_KEY", "")
CLIENT_SECRET = os.getenv("DS_CLIENT_SECRET", "")
REFRESH_TOKEN = os.getenv("DS_REFRESH_TOKEN", "")
ACCOUNT_ID = os.getenv("DS_ACCOUNT_ID", "")

START_DATE = os.getenv("DS_START_DATE", "2025-05-01")

OUTPUT_FILE = Path(os.getenv("DS_OUTPUT_FILE", "docusign_envelope_usage.xlsx"))

# "created" -> envelopes CREATED in the month (use this for usage counting)
# "changed" -> envelopes whose status CHANGED in the month (DocuSign default)
FROM_TO_STATUS = os.getenv("DS_FROM_TO_STATUS", "created")

# fast : one count=1 call per status, read totalSetSize   (~11 calls/month)
# full : page every envelope and tally locally, deduped   (~90 calls/month)
COUNT_MODE = os.getenv("DS_COUNT_MODE", "fast").lower()

# Month boundaries are built in this offset then converted to UTC for the API.
# Set it to your DocuSign account's timezone to line up with the Admin dashboard.
TZ_OFFSET = os.getenv("DS_TZ_OFFSET", "+00:00")

REBUILD = os.getenv("DS_REBUILD", "") == "1"

HTTP_TIMEOUT = int(os.getenv("DS_TIMEOUT", "180"))
ENVELOPES_PAGE_SIZE = 100

AUTH_HOST = "account-d.docusign.com" if ENVIRONMENT == "demo" else "account.docusign.com"

MONTHLY_SHEET = "MonthlyUsage"
TOTALS_SHEET = "TotalsToDate"

# ---------------------------------------------------------------------------
# STATUS MODEL
# ---------------------------------------------------------------------------
#
# Left  = the status string the REST API accepts and returns.
# Right = the column name written to Excel.
#
# The Admin dashboard's eight buckets do not map 1:1 onto API statuses. Its
# "Expired" is the API's "timedout", its "Corrected" is "correct", and its
# "In Progress" is an aggregate. Raw API statuses are stored so nothing is lost,
# and dashboard-shaped rollups are derived from them further down.

API_STATUSES = [
    ("completed", "Completed"),
    ("declined", "Declined"),
    ("voided", "Voided"),
    ("deleted", "Deleted"),
    ("timedout", "Expired"),
    ("correct", "Corrected"),
    ("created", "Created"),
    ("sent", "Sent"),
    ("delivered", "Delivered"),
    ("signed", "Signed"),
]

STATUS_COLUMNS = [col for _api, col in API_STATUSES]

DATE_COLUMNS = [
    "MonthStart",      # real date, 1st of month -- join key for a Power BI date table
    "MonthEnd",        # real date, last day actually counted
    "Year",            # 2025
    "MonthNumber",     # 1-12
    "MonthName",       # May
    "MonthYear",       # May 2025      <- use as the axis label
    "YearMonthKey",    # 202505        <- sort MonthYear by this, or Power BI sorts A-Z
    "DaysInMonth",
    "DaysCounted",
    "IsComplete",
]

DERIVED_COLUMNS = [
    "InProgress",      # Created + Sent + Delivered + Signed
    "TotalEnvelopes",  # unfiltered count straight from DocuSign
    "SumOfStatuses",   # the status columns added up
    "CountVariance",   # TotalEnvelopes - SumOfStatuses; should be 0
]

MONTHLY_HEADER = DATE_COLUMNS + STATUS_COLUMNS + DERIVED_COLUMNS + ["FetchedUTC"]


# ---------------------------------------------------------------------------
# AUTH
# ---------------------------------------------------------------------------


def get_access_token():
    if ACCESS_TOKEN:
        return ACCESS_TOKEN

    if not (INTEGRATION_KEY and CLIENT_SECRET and REFRESH_TOKEN):
        sys.exit(
            "No credentials. Either:\n"
            "  set DS_ACCESS_TOKEN=<token>\n"
            "or, to generate one:\n"
            "  set DS_INTEGRATION_KEY=<yml username>\n"
            "  set DS_CLIENT_SECRET=<yml password>\n"
            "  set DS_REFRESH_TOKEN=<yml refresh_token>"
        )

    print("No DS_ACCESS_TOKEN set -- generating one from the refresh token.")
    print("(This rotates the refresh token.)\n")

    basic = base64.b64encode(f"{INTEGRATION_KEY}:{CLIENT_SECRET}".encode()).decode()
    resp = requests.post(
        f"https://{AUTH_HOST}/oauth/token",
        headers={"Authorization": f"Basic {basic}"},
        data={"grant_type": "refresh_token", "refresh_token": REFRESH_TOKEN},
        timeout=30,
    )

    if resp.status_code != 200:
        sys.exit(
            f"Token request failed ({resp.status_code}). Common causes: the "
            "refresh token was already rotated by someone else, it has expired "
            "(~30 days), or the integration key / secret don't match the app "
            f"that issued it.\n\nDocuSign said: {resp.text}"
        )

    data = resp.json()
    new_refresh = data.get("refresh_token")

    try:
        lines = [f'set DS_ACCESS_TOKEN={data["access_token"]}']
        if new_refresh and new_refresh != REFRESH_TOKEN:
            lines += ["", f"set DS_REFRESH_TOKEN={new_refresh}"]
        Path("docusign_tokens.txt").write_text("\n".join(lines), encoding="utf-8")
        wrote_file = True
    except OSError:
        wrote_file = False

    if new_refresh and new_refresh != REFRESH_TOKEN:
        print("=" * 70)
        print("NEW REFRESH TOKEN -- save this, the old one may now be dead:\n")
        print(new_refresh)
        print("\nTell anyone else sharing these credentials.")
        print("=" * 70 + "\n")

    print("=" * 70)
    print("ACCESS TOKEN (valid ~8 hours). To re-run without rotating again:\n")
    print(f'set DS_ACCESS_TOKEN={data["access_token"]}')
    print("=" * 70)
    if wrote_file:
        print("\nAlso saved to docusign_tokens.txt. Delete it when you're done.\n")

    return data["access_token"]


def get_account_context(token):
    resp = requests.get(
        f"https://{AUTH_HOST}/oauth/userinfo",
        headers={"Authorization": f"Bearer {token}"},
        timeout=30,
    )

    if resp.status_code == 401:
        sys.exit(
            "Token rejected (401). If you supplied DS_ACCESS_TOKEN it has most "
            "likely expired -- these last about 8 hours. Either paste a fresh "
            "one, or unset DS_ACCESS_TOKEN and let the script generate one.\n\n"
            f"If the token is definitely fresh, check DS_ENV: it's '{ENVIRONMENT}', "
            f"so the token is going to {AUTH_HOST}."
        )

    resp.raise_for_status()
    accounts = resp.json()["accounts"]

    if ACCOUNT_ID:
        account = next((a for a in accounts if a["account_id"] == ACCOUNT_ID), None)
        if account is None:
            sys.exit(f"Account {ACCOUNT_ID} not visible to this user.")
    else:
        account = next((a for a in accounts if a.get("is_default")), accounts[0])

    base_path = f"{account['base_uri']}/restapi/v2.1/accounts/{account['account_id']}"
    return base_path, account["account_id"], account.get("account_name", "")


# ---------------------------------------------------------------------------
# HTTP
# ---------------------------------------------------------------------------


def api_get(url, token, params=None, max_retries=5):
    for attempt in range(max_retries):
        try:
            resp = requests.get(
                url,
                headers={"Authorization": f"Bearer {token}", "Accept": "application/json"},
                params=params,
                timeout=HTTP_TIMEOUT,
            )
        except (requests.Timeout, requests.ConnectionError) as e:
            if attempt == max_retries - 1:
                sys.exit(
                    f"Request kept timing out after {max_retries} attempts:\n  {url}\n\n"
                    "Try:  set DS_TIMEOUT=300\n"
                    f"Underlying error: {e}"
                )
            wait = 5 * (attempt + 1)
            print(f"    timed out, retrying in {wait}s ...", flush=True)
            time.sleep(wait)
            continue

        if resp.status_code == 429:
            reset = resp.headers.get("X-RateLimit-Reset")
            wait = max(int(reset) - int(time.time()), 1) if reset else 2 ** attempt * 30
            print(f"    rate limited, sleeping {wait}s ...", flush=True)
            time.sleep(min(wait, 3700))
            continue

        if resp.status_code >= 500:
            time.sleep(2 ** attempt)
            continue

        if 400 <= resp.status_code < 500:
            try:
                err = resp.json()
                detail = f"{err.get('errorCode', '')}: {err.get('message', resp.text)}"
            except ValueError:
                detail = resp.text
            sys.exit(f"DocuSign rejected the request ({resp.status_code}).\n\n  {detail}\n\nURL: {resp.url}")

        resp.raise_for_status()

        remaining = resp.headers.get("X-RateLimit-Remaining")
        if remaining is not None and int(remaining) < 50:
            print(f"    warning: only {remaining} API calls left this hour", flush=True)

        return resp.json()

    raise RuntimeError(f"Giving up on {url} after {max_retries} attempts")


# ---------------------------------------------------------------------------
# MONTH WINDOWS
# ---------------------------------------------------------------------------


def parse_offset(text):
    """'+05:30' -> timedelta. Accepts '+0530', '-05:00', '+00:00'."""
    text = text.strip()
    if not text or text.upper() == "Z":
        return timedelta(0)
    sign = -1 if text[0] == "-" else 1
    body = text.lstrip("+-").replace(":", "")
    if len(body) != 4:
        sys.exit(f"DS_TZ_OFFSET '{text}' is not in +HH:MM form.")
    return sign * timedelta(hours=int(body[:2]), minutes=int(body[2:]))


OFFSET = parse_offset(TZ_OFFSET)


def to_utc_string(naive_local):
    """Local wall-clock datetime -> the UTC instant DocuSign should be given."""
    return (naive_local - OFFSET).strftime("%Y-%m-%dT%H:%M:%SZ")


def iter_months(start_date_str):
    """
    Yield one descriptor per month from the start month through the current
    month inclusive. The current month stops at today (23:59:59 local).
    """
    start = datetime.strptime(start_date_str, "%Y-%m-%d")
    today = datetime.now(timezone.utc) + OFFSET  # today in account-local terms
    today = today.replace(tzinfo=None)

    year, month = start.year, start.month
    while (year, month) <= (today.year, today.month):
        days_in_month = calendar.monthrange(year, month)[1]
        is_current = (year, month) == (today.year, today.month)
        last_day = today.day if is_current else days_in_month

        yield {
            "year": year,
            "month": month,
            "days_in_month": days_in_month,
            "days_counted": last_day,
            "is_complete": not is_current,
            "window_start": datetime(year, month, 1, 0, 0, 0),
            "window_end": datetime(year, month, last_day, 23, 59, 59),
        }

        month += 1
        if month > 12:
            month, year = 1, year + 1


# ---------------------------------------------------------------------------
# COUNTING
# ---------------------------------------------------------------------------


def count_fast(base_path, token, window_start, window_end):
    """
    One request per status with count=1, reading totalSetSize off the response
    rather than downloading the envelopes. Plus one unfiltered request for the
    true total, which is what CountVariance is checked against.
    """
    base_params = {
        "from_date": to_utc_string(window_start),
        "to_date": to_utc_string(window_end),
        "from_to_status": FROM_TO_STATUS,
        "count": 1,
        "start_position": 0,
    }

    counts = {}
    for api_status, column in API_STATUSES:
        data = api_get(f"{base_path}/envelopes", token, params={**base_params, "status": api_status})
        counts[column] = int(data.get("totalSetSize", 0) or 0)

    data = api_get(f"{base_path}/envelopes", token, params=base_params)
    total = int(data.get("totalSetSize", 0) or 0)

    return counts, total


def count_full(base_path, token, window_start, window_end):
    """
    Page through every envelope in the window and tally locally, deduping on
    envelopeId. Slower and heavier, but it is ground truth -- use it to prove
    the fast mode agrees, then go back to fast.
    """
    status_by_api = dict(API_STATUSES)
    counts = {col: 0 for col in STATUS_COLUMNS}
    seen = set()
    unknown = {}

    start_pos = 0
    while True:
        data = api_get(
            f"{base_path}/envelopes",
            token,
            params={
                "from_date": to_utc_string(window_start),
                "to_date": to_utc_string(window_end),
                "from_to_status": FROM_TO_STATUS,
                "count": ENVELOPES_PAGE_SIZE,
                "start_position": start_pos,
                "order_by": "created",
                "order": "asc",
            },
        )

        batch = data.get("envelopes", []) or []
        if not batch:
            break

        for env in batch:
            eid = env.get("envelopeId")
            if eid in seen:
                continue          # window boundaries are inclusive on both ends
            seen.add(eid)
            raw = (env.get("status") or "").lower()
            column = status_by_api.get(raw)
            if column:
                counts[column] += 1
            else:
                unknown[raw] = unknown.get(raw, 0) + 1

        start_pos += len(batch)
        total = int(data.get("totalSetSize", 0) or 0)
        if total and start_pos >= total:
            break

    if unknown:
        print(f"    statuses seen but not modelled: {unknown}", flush=True)

    return counts, len(seen)


def build_row(descriptor, counts, total):
    y, m = descriptor["year"], descriptor["month"]
    in_progress = sum(counts.get(c, 0) for c in ("Created", "Sent", "Delivered", "Signed"))
    sum_statuses = sum(counts.get(c, 0) for c in STATUS_COLUMNS)

    row = {
        "MonthStart": datetime(y, m, 1).date(),
        "MonthEnd": datetime(y, m, descriptor["days_counted"]).date(),
        "Year": y,
        "MonthNumber": m,
        "MonthName": calendar.month_name[m],
        "MonthYear": f"{calendar.month_abbr[m]} {y}",
        "YearMonthKey": y * 100 + m,
        "DaysInMonth": descriptor["days_in_month"],
        "DaysCounted": descriptor["days_counted"],
        "IsComplete": descriptor["is_complete"],
        "InProgress": in_progress,
        "TotalEnvelopes": total,
        "SumOfStatuses": sum_statuses,
        "CountVariance": total - sum_statuses,
        "FetchedUTC": datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M:%S"),
    }
    row.update({c: counts.get(c, 0) for c in STATUS_COLUMNS})
    return row


# ---------------------------------------------------------------------------
# WORKBOOK I/O
# ---------------------------------------------------------------------------


def read_existing():
    """Return {(year, month): row_dict} from the workbook, if one exists."""
    if REBUILD or not OUTPUT_FILE.exists():
        return {}

    try:
        wb = load_workbook(OUTPUT_FILE)
    except Exception as e:
        print(f"  could not open {OUTPUT_FILE} ({e}); starting a fresh workbook.")
        return {}

    if MONTHLY_SHEET not in wb.sheetnames:
        return {}

    ws = wb[MONTHLY_SHEET]
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return {}

    header = [str(h) if h is not None else "" for h in rows[0]]
    existing = {}

    for values in rows[1:]:
        record = dict(zip(header, values))
        try:
            key = (int(record["Year"]), int(record["MonthNumber"]))
        except (KeyError, TypeError, ValueError):
            continue
        existing[key] = record

    wb.close()
    return existing


def normalise(record):
    """Coerce a row read back out of Excel into the shape we write."""
    out = {}
    for col in MONTHLY_HEADER:
        value = record.get(col)
        if col in ("MonthStart", "MonthEnd") and isinstance(value, datetime):
            value = value.date()
        if col == "IsComplete":
            value = bool(value) and str(value).upper() != "FALSE"
        out[col] = value
    return out


def write_workbook(rows):
    rows = sorted(rows, key=lambda r: r["YearMonthKey"])

    wb = Workbook()

    ws = wb.active
    ws.title = MONTHLY_SHEET
    ws.append(MONTHLY_HEADER)
    for row in rows:
        ws.append([row.get(c) for c in MONTHLY_HEADER])

    for col in ("A", "B"):
        for cell in ws[col][1:]:
            cell.number_format = "yyyy-mm-dd"
    ws.freeze_panes = "A2"

    # ---- totals sheet: one row, rebuilt from scratch every run ----
    totals = wb.create_sheet(TOTALS_SHEET)

    complete_rows = [r for r in rows if r.get("IsComplete")]
    header = (
        ["LastUpdatedUTC", "RangeStart", "RangeEnd", "MonthsCovered",
         "CompleteMonths", "PartialMonths", "TotalEnvelopes",
         "TotalEnvelopesCompleteMonthsOnly", "AvgEnvelopesPerCompleteMonth"]
        + STATUS_COLUMNS
        + ["InProgress"]
    )
    totals.append(header)

    grand_total = sum(int(r.get("TotalEnvelopes") or 0) for r in rows)
    complete_total = sum(int(r.get("TotalEnvelopes") or 0) for r in complete_rows)

    record = {
        "LastUpdatedUTC": datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M:%S"),
        "RangeStart": rows[0]["MonthStart"] if rows else None,
        "RangeEnd": rows[-1]["MonthEnd"] if rows else None,
        "MonthsCovered": len(rows),
        "CompleteMonths": len(complete_rows),
        "PartialMonths": len(rows) - len(complete_rows),
        "TotalEnvelopes": grand_total,
        "TotalEnvelopesCompleteMonthsOnly": complete_total,
        "AvgEnvelopesPerCompleteMonth": (
            round(complete_total / len(complete_rows), 1) if complete_rows else 0
        ),
        "InProgress": sum(int(r.get("InProgress") or 0) for r in rows),
    }
    for col in STATUS_COLUMNS:
        record[col] = sum(int(r.get(col) or 0) for r in rows)

    totals.append([record.get(c) for c in header])
    for col in ("B", "C"):
        for cell in totals[col][1:]:
            cell.number_format = "yyyy-mm-dd"

    wb.save(OUTPUT_FILE)
    return record


# ---------------------------------------------------------------------------
# MAIN
# ---------------------------------------------------------------------------


def main():
    if COUNT_MODE not in ("fast", "full"):
        sys.exit("DS_COUNT_MODE must be 'fast' or 'full'.")

    print("Authenticating ...")
    token = get_access_token()
    base_path, account_id, account_name = get_account_context(token)
    print(f"  account: {account_name} ({account_id})")
    print(f"  mode:    {COUNT_MODE}   basis: from_to_status={FROM_TO_STATUS}   tz: {TZ_OFFSET}")

    existing = read_existing()
    if existing:
        print(f"  found {len(existing)} month(s) already in {OUTPUT_FILE.name}")
    elif REBUILD:
        print("  DS_REBUILD=1 -- refetching every month")

    rows, fetched, skipped = [], 0, 0

    for descriptor in iter_months(START_DATE):
        key = (descriptor["year"], descriptor["month"])
        label = f"{calendar.month_abbr[descriptor['month']]} {descriptor['year']}"
        prior = existing.get(key)

        # Reuse only a stored row that was complete when it was captured.
        if prior is not None and descriptor["is_complete"]:
            was_complete = bool(prior.get("IsComplete")) and str(prior.get("IsComplete")).upper() != "FALSE"
            if was_complete:
                rows.append(normalise(prior))
                skipped += 1
                continue
            print(f"  {label}: stored row was partial -- refetching")

        reason = "current month" if not descriptor["is_complete"] else "new"
        print(f"  {label}: fetching ({reason}, days 1-{descriptor['days_counted']}) ...", flush=True)

        counter = count_fast if COUNT_MODE == "fast" else count_full
        counts, total = counter(base_path, token, descriptor["window_start"], descriptor["window_end"])

        row = build_row(descriptor, counts, total)
        rows.append(row)
        fetched += 1

        if row["CountVariance"]:
            print(
                f"    note: variance {row['CountVariance']:+d} "
                f"(total {row['TotalEnvelopes']} vs statuses {row['SumOfStatuses']})"
            )

    if not rows:
        sys.exit("No months in range -- check DS_START_DATE.")

    summary = write_workbook(rows)

    print(f"\nWrote {OUTPUT_FILE}")
    print(f"  {len(rows)} months  ({fetched} fetched, {skipped} reused from file)")
    print(f"  total envelopes to date: {summary['TotalEnvelopes']:,}")
    print(f"  complete months only:    {summary['TotalEnvelopesCompleteMonthsOnly']:,}")


if __name__ == "__main__":
    main()
