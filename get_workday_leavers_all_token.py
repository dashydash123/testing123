"""
Scheduled pull of the full Workday schema from the Databricks SQL warehouse.
Every table/view in the schema goes into its own sheet of one Excel workbook.

Setup:
    pip install databricks-sql-connector databricks-sdk pandas pyarrow openpyxl

Auth:
    Paste your Databricks personal access token into TOKEN below.
    (If TOKEN is left blank, falls back to env vars DATABRICKS_CLIENT_ID /
    DATABRICKS_CLIENT_SECRET or DATABRICKS_TOKEN.)
"""
import os
import re
import sys
from datetime import date
import pandas as pd
from databricks import sql
from databricks.sdk.core import Config, oauth_service_principal

# >>> Paste your Databricks access token here <<<
TOKEN = "PASTE_YOUR_TOKEN_HERE"

HOST = "adb-3060038306817150.10.azuredatabricks.net"
HTTP_PATH = "/sql/1.0/warehouses/13861ac29ff00f3e"
CATALOG = "brewdat_uc_people_prod"
SCHEMA = "gld_ghq_people_workday_tr"

EXCEL_MAX_ROWS = 1_048_575  # Excel limit minus header row
ILLEGAL_CHARS = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


def connect():
    client_id = os.getenv("DATABRICKS_CLIENT_ID")
    client_secret = os.getenv("DATABRICKS_CLIENT_SECRET")
    token = TOKEN if TOKEN and TOKEN != "PASTE_YOUR_TOKEN_HERE" \
        else os.getenv("DATABRICKS_TOKEN")

    if token:
        return sql.connect(server_hostname=HOST, http_path=HTTP_PATH,
                           access_token=token)

    if client_id and client_secret:
        def credentials_provider():
            cfg = Config(host=f"https://{HOST}", client_id=client_id,
                         client_secret=client_secret)
            return oauth_service_principal(cfg)
        return sql.connect(server_hostname=HOST, http_path=HTTP_PATH,
                           credentials_provider=credentials_provider)
    sys.exit("No credentials: paste your token into TOKEN at the top of the script")


def list_objects(cur):
    """All tables and views in the schema that this identity can read."""
    cur.execute(f"""
        SELECT table_name, table_type
        FROM {CATALOG}.information_schema.tables
        WHERE table_schema = '{SCHEMA}'
        ORDER BY table_name
    """)
    return cur.fetchall()


def sheet_name(name, used):
    """Excel sheet names: max 31 chars, no []:*?/\\ , must be unique."""
    base = re.sub(r"[\[\]:*?/\\]", "_", name)[:31]
    candidate, n = base, 1
    while candidate.lower() in used:
        suffix = f"~{n}"
        candidate = base[:31 - len(suffix)] + suffix
        n += 1
    used.add(candidate.lower())
    return candidate


def clean_for_excel(df):
    """Strip timezones and control characters that Excel/openpyxl reject."""
    for col in df.columns:
        if isinstance(df[col].dtype, pd.DatetimeTZDtype):
            df[col] = df[col].dt.tz_localize(None)
        elif pd.api.types.is_object_dtype(df[col]) or pd.api.types.is_string_dtype(df[col]):
            df[col] = df[col].map(
                lambda v: ILLEGAL_CHARS.sub("", v) if isinstance(v, str) else v)
    return df


def main():
    out = f"workday_{SCHEMA}_{date.today():%Y%m%d}.xlsx"
    index_rows, used, failures = [], {"index"}, 0

    with connect() as conn, conn.cursor() as cur, \
            pd.ExcelWriter(out, engine="openpyxl") as writer:
        objects = list_objects(cur)
        if not objects:
            sys.exit(f"No readable tables/views found in {CATALOG}.{SCHEMA}")

        for table_name, table_type in objects:
            full = f"{CATALOG}.{SCHEMA}.{table_name}"
            try:
                cur.execute(f"SELECT * FROM {full}")
                df = clean_for_excel(cur.fetchall_arrow().to_pandas())

                # Split across extra sheets if over Excel's row limit
                chunks = range(0, max(len(df), 1), EXCEL_MAX_ROWS)
                sheets = []
                for i, start in enumerate(chunks):
                    label = table_name if i == 0 else f"{table_name}_{i + 1}"
                    name = sheet_name(label, used)
                    df.iloc[start:start + EXCEL_MAX_ROWS].to_excel(
                        writer, sheet_name=name, index=False)
                    sheets.append(name)

                index_rows.append([", ".join(sheets), full, table_type,
                                   len(df), "OK"])
                print(f"{full}: {len(df)} rows")
            except Exception as e:
                failures += 1
                index_rows.append(["", full, table_type, None, f"FAILED: {e}"])
                print(f"{full}: FAILED - {e}")

        pd.DataFrame(index_rows, columns=["Sheet", "Source object", "Type",
                                          "Rows", "Status"]
                     ).to_excel(writer, sheet_name="Index", index=False)

    # Put the Index sheet first
    from openpyxl import load_workbook
    wb = load_workbook(out)
    wb.move_sheet("Index", offset=-(len(wb.sheetnames) - 1))
    wb.save(out)

    print(f"Saved {len(objects)} objects to {out}")
    if failures:
        sys.exit(f"{failures} object(s) failed - see Index sheet")


if __name__ == "__main__":
    try:
        main()
    except SystemExit:
        raise
    except Exception as e:
        # Non-zero exit so the scheduler flags the failure
        sys.exit(f"Refresh failed: {e}")
